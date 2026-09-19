#!/usr/bin/env python3
"""Find broker export files that were downloaded but never imported, and import them.

Two separate mechanisms quietly dropped files on the floor:

1. The download folder moved. Until 2026-09-10 everything landed in the
   `~/Downloads` root; after that each account got its own
   `~/Downloads/energrex_<nnn>/`. Anything Chrome kept writing to the old
   location stopped being seen.

2. The watchdog only ever imports a CSV **on the day it was downloaded**
   (`account_monitor._CsvHandler._handle`: `if mdate != datetime.date.today():
   return`). A file that arrived while the app was closed is skipped that day
   and then skipped forever -- nothing ever revisits it. Missing transactions
   break FIFO matching and leave gaps in the performance history.

This script scans the likely locations, reports what it finds with each file's
real date, and (with --import) feeds the CSVs through the normal importer,
bypassing the same-day rule.

Read-only by default:

    python scripts/recover_broker_downloads.py                  # report only
    python scripts/recover_broker_downloads.py --import         # also import CSVs
    python scripts/recover_broker_downloads.py --move           # file xlsx into the account folder
    python scripts/recover_broker_downloads.py --extra-dir "D:/somewhere"

What each file type can and cannot restore:

  export*.csv    transactions and/or positions -- importable, fills history gaps
  positions.xlsx holdings and market values at that moment -- no cash, no equity,
                 so it CANNOT rebuild the NAV curve on its own. The `daily_nav`
                 curve is written by `save_balance()` during a live sync; a gap
                 there needs either a successful sync or month-end equity from
                 the statement PDFs (see energrex-risk-monitor's
                 scripts/extract_account_equity_history.py).
"""

from __future__ import annotations

import argparse
import datetime as dt
import pathlib
import sys
import zipfile

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

CSV_PATTERNS = ("export*.csv", "Export*.csv", "EXPORT*.csv")
XLSX_PATTERNS = ("*.xlsx",)


def candidate_dirs(extra: list[str] | None = None) -> list[pathlib.Path]:
    """Every place a Firstrade download has plausibly landed."""
    home = pathlib.Path.home()
    dirs = [home / "Downloads"]
    downloads = home / "Downloads"
    if downloads.exists():
        # per-account subfolders introduced 2026-09-10
        dirs.extend(sorted(p for p in downloads.glob("energrex_*") if p.is_dir()))
    dirs.append(home / "Desktop")
    dirs.append(ROOT / "imports")
    for name in (extra or []):
        dirs.append(pathlib.Path(name).expanduser())
    seen: set[pathlib.Path] = set()
    unique: list[pathlib.Path] = []
    for d in dirs:
        try:
            resolved = d.resolve()
        except OSError:
            continue
        if resolved in seen:
            continue
        seen.add(resolved)
        unique.append(d)
    return unique


def xlsx_internal_date(path: pathlib.Path) -> dt.datetime | None:
    """When the workbook was actually generated.

    The file's mtime says when it was copied or moved; the zip entry timestamps
    say when the exporter wrote it. For diagnosing "which day is this snapshot
    really from" only the latter can be trusted.
    """
    try:
        with zipfile.ZipFile(path) as zf:
            stamps = [i.date_time for i in zf.infolist() if i.date_time[0] > 1980]
        if stamps:
            return dt.datetime(*max(stamps))
    except (zipfile.BadZipFile, OSError, ValueError):
        pass
    return None


def describe_xlsx(path: pathlib.Path) -> str:
    try:
        import openpyxl
    except ImportError:
        return "（未安装 openpyxl，无法读取内容）"
    try:
        # 不用 read_only：那个模式下 max_row 可能是 None（取决于写入方有没有
        # 写 dimension 记录，Firstrade 的 AG Grid 导出就没写）。
        wb = openpyxl.load_workbook(str(path), read_only=False, data_only=True)
        try:
            parts = []
            for name in wb.sheetnames:
                rows = wb[name].max_row
                count = max((rows or 1) - 1, 0)  # 减表头
                parts.append(f"{name}:{count}行")
        finally:
            wb.close()
        return "  ".join(parts)
    except Exception as exc:  # noqa: BLE001 - diagnostics only
        return f"（读取失败: {exc}）"


def scan(dirs: list[pathlib.Path]) -> tuple[list[dict], list[dict]]:
    csvs: list[dict] = []
    xlsxs: list[dict] = []
    for directory in dirs:
        if not directory.exists():
            continue
        for pattern in CSV_PATTERNS:
            for path in directory.glob(pattern):
                csvs.append({
                    "path": path,
                    "when": dt.datetime.fromtimestamp(path.stat().st_mtime),
                    "size": path.stat().st_size,
                })
        for pattern in XLSX_PATTERNS:
            for path in directory.glob(pattern):
                if path.name.startswith("~$"):  # Excel lock file
                    continue
                xlsxs.append({
                    "path": path,
                    "when": xlsx_internal_date(path)
                            or dt.datetime.fromtimestamp(path.stat().st_mtime),
                    "size": path.stat().st_size,
                })
    dedupe = {item["path"].resolve(): item for item in csvs}
    csvs = sorted(dedupe.values(), key=lambda i: i["when"])
    dedupe = {item["path"].resolve(): item for item in xlsxs}
    xlsxs = sorted(dedupe.values(), key=lambda i: i["when"])
    return csvs, xlsxs


def import_csvs(found: list[dict], acct_id: str) -> None:
    """Run each CSV through the normal importer, ignoring the same-day rule."""
    import logging

    from account.importers import process_csv_file
    from account.repository import save_balance, save_positions, save_transactions

    logging.basicConfig(level=logging.INFO, format="  %(message)s")
    logger = logging.getLogger("recover")

    latest_csv = ROOT / "data" / "latest_import.csv"
    latest_csv.parent.mkdir(parents=True, exist_ok=True)

    for item in found:
        path = item["path"]
        print(f"\n导入 {path}  ({item['when']:%Y-%m-%d %H:%M})")
        try:
            result = process_csv_file(
                path,
                acct_id=acct_id,
                latest_csv=latest_csv,
                save_positions=save_positions,
                save_transactions=save_transactions,
                save_balance=save_balance,
                logger=logger,
            )
            print(f"  → {result}")
        except Exception as exc:  # noqa: BLE001 - one bad file must not stop the rest
            print(f"  → 失败: {exc}")


def move_xlsx(found: list[dict], acct_id: str) -> None:
    from account.accounts import account_download_dir

    target_dir = account_download_dir(acct_id)
    for item in found:
        path = item["path"]
        if path.parent.resolve() == target_dir.resolve():
            continue
        target = target_dir / f"{item['when']:%Y%m%d_%H%M}_{path.name}"
        try:
            path.replace(target)
            print(f"  移动 {path}  →  {target}")
        except OSError as exc:
            print(f"  移动失败 {path}: {exc}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--account", default="account_1")
    parser.add_argument("--extra-dir", action="append", default=[],
                        help="额外扫描目录，可重复")
    parser.add_argument("--import", dest="do_import", action="store_true",
                        help="把找到的 export*.csv 导入数据库（绕过“只认当天”限制）")
    parser.add_argument("--move", action="store_true",
                        help="把散落的 positions xlsx 归档到账户下载目录")
    args = parser.parse_args()

    dirs = candidate_dirs(args.extra_dir)
    print("扫描目录：")
    for d in dirs:
        print(f"  {'✓' if d.exists() else '✗'} {d}")

    csvs, xlsxs = scan(dirs)

    print(f"\n找到 {len(csvs)} 个 export*.csv（成交/持仓导出，可导入）：")
    for item in csvs:
        print(f"  {item['when']:%Y-%m-%d %H:%M}  {item['size']:>9,}B  {item['path']}")
    if not csvs:
        print("  （无）")

    print(f"\n找到 {len(xlsxs)} 个 xlsx（持仓快照；不含现金/净值，补不了收益曲线）：")
    for item in xlsxs:
        print(f"  {item['when']:%Y-%m-%d %H:%M}  {item['size']:>9,}B  {item['path']}")
        print(f"      {describe_xlsx(item['path'])}")
    if not xlsxs:
        print("  （无）")

    if args.move and xlsxs:
        print("\n归档 xlsx：")
        move_xlsx(xlsxs, args.account)

    if args.do_import:
        if csvs:
            print("\n开始导入 CSV（按日期从旧到新）：")
            import_csvs(csvs, args.account)
        else:
            print("\n没有可导入的 CSV。")
    elif csvs:
        print("\n以上 CSV 尚未导入。确认无误后加 --import 执行导入。")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
