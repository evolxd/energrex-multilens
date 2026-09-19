"""Parse the stock sheet of a Firstrade positions export into `positions` rows.

Why this file exists: `account_monitor._parse_positions_xlsx()` only ever read
the options sheet, dropping every stock row on the floor
(`if not occ: continue  # 跳过非期权行（股票等）`). Its caller then decided
whether to fall back to the stock-scraping JS path with `if not new_rows:` --
and for any account holding options that condition is never true. Stocks
therefore stopped being refreshed entirely, frozen at whatever the last
`export*.csv` import happened to write.

That is not a cosmetic gap. Gate ③ hard limits and gate ④ pre-trade checks both
size single-stock and supply-chain concentration off this table. Measured on the
real account on 2026-09-19: the dashboard reported a 15.4% largest position with
"+2.6% headroom" while the true largest was ETHU at 19.6%, already through the
18% cap -- the limit system was clearing trades it existed to block.

Kept out of `account_monitor.py` deliberately: importing that module executes a
whole Streamlit page, so pure parsing logic living there cannot be unit tested.
Same split the architecture review already prescribes for `account/options.py`.
"""

from __future__ import annotations

import pathlib
from typing import Iterable

from account.options import parse_occ_sym

# Firstrade localises its export. Columns and sheet names are matched in both
# languages rather than by position, because the column order has changed
# between exports before.
_STOCK_SHEET_HINTS = ("股票", "stock", "etf")
_OPTION_SHEET_HINTS = ("期权", "option")

_HEADER_KEYS = ("代号", "Symbol")


def _is_stock_sheet(name: str) -> bool:
    lower = name.lower()
    if any(hint in name or hint in lower for hint in _OPTION_SHEET_HINTS):
        return False
    return any(hint in name or hint in lower for hint in _STOCK_SHEET_HINTS)


def _to_number(value) -> float | None:
    """Firstrade mixes types within one column: Market Value arrives as a float,
    Total Cost as '1,036.00', % Gain/Loss as '-4.11%'. All three must land as
    numbers or concentration maths silently sees None.
    """
    if value is None:
        return None
    try:
        return float(
            str(value).replace(",", "").replace("$", "").replace("%", "").strip()
        )
    except (ValueError, TypeError):
        return None


def parse_stocks_xlsx(xlsx_path: pathlib.Path) -> list[dict]:
    """Return `positions`-table row dicts for every stock line in the export.

    Returns an empty list rather than raising when the sheet or header is
    missing -- the caller treats that as "keep the previous snapshot", which is
    the safe reading of a malformed download.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(str(xlsx_path), data_only=True, read_only=True)
    try:
        sheet = next(
            (workbook[name] for name in workbook.sheetnames if _is_stock_sheet(name)),
            None,
        )
        if sheet is None:
            return []
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    header: list[str] | None = None
    header_index = 0
    for index, row in enumerate(rows):
        values = [str(cell or "").strip() for cell in row]
        if any(key in values for key in _HEADER_KEYS):
            header, header_index = values, index
            break
    if header is None:
        return []

    column = {name: i for i, name in enumerate(header)}

    def cell(row: tuple, *names: str):
        for name in names:
            index = column.get(name)
            if index is not None and index < len(row):
                value = row[index]
                if value is not None and str(value).strip():
                    return value
        return None

    results: list[dict] = []
    seen: set[str] = set()
    for row in rows[header_index + 1:]:
        raw_symbol = cell(row, "代号", "Symbol")
        if not raw_symbol:
            continue
        symbol = str(raw_symbol).strip().upper().replace(" ", "")
        if not symbol or symbol in seen:
            continue
        if parse_occ_sym(symbol):
            continue  # an option contract that wandered into the stock sheet

        quantity = _to_number(cell(row, "数量", "Quantity", "Qty"))
        if quantity is None:
            continue

        description = cell(row, "Description", "详细说明", "說明")
        results.append({
            "symbol": symbol,
            "position_type": "stock",
            "quantity": quantity,
            "cost_basis": _to_number(cell(row, "Total Cost", "总成本", "成本")),
            "market_value": _to_number(cell(row, "Market Value", "市值", "MktVal")),
            "unrealized_pnl": _to_number(
                cell(row, "$ Gain/Loss", "益损 $", "益损$", "Gain/Loss $")),
            "unrealized_pnl_pct": _to_number(
                cell(row, "% Gain/Loss", "益损 %", "益损%", "Gain/Loss %")),
            "description": str(description).strip() if description else None,
        })
        seen.add(symbol)

    return results


def parse_options_xlsx(xlsx_path: pathlib.Path) -> list[dict]:
    """Return `options_positions` row dicts for every option line in the export.

    Ported here from `account_monitor._parse_positions_xlsx` so the offline
    importer can reach it: that module executes a Streamlit page on import, so
    a command-line tool could not use the parser while it lived there. The
    account_monitor entry point now delegates to this function, keeping one
    implementation.

    A negative quantity means a short leg.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(str(xlsx_path), data_only=True, read_only=True)
    try:
        sheet = next(
            (workbook[name] for name in workbook.sheetnames
             if any(hint in name or hint in name.lower() for hint in _OPTION_SHEET_HINTS)),
            None,
        )
        if sheet is None:
            return []
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    header: list[str] | None = None
    header_index = 0
    for index, row in enumerate(rows):
        values = [str(cell or "").strip() for cell in row]
        if any(key in values for key in _HEADER_KEYS):
            header, header_index = values, index
            break
    if header is None:
        return []

    column = {name: i for i, name in enumerate(header)}

    def cell(row: tuple, *names: str):
        for name in names:
            index = column.get(name)
            if index is not None and index < len(row):
                value = row[index]
                if value is not None and str(value).strip():
                    return value
        return None

    results: list[dict] = []
    seen: set[str] = set()
    for row in rows[header_index + 1:]:
        raw_symbol = cell(row, "代号", "Symbol")
        if not raw_symbol:
            continue
        symbol = str(raw_symbol).strip().upper().replace(" ", "")
        if not symbol or symbol in seen:
            continue
        occ = parse_occ_sym(symbol)
        if not occ:
            continue  # a stock row that wandered into the options sheet

        quantity = _to_number(cell(row, "数量", "Quantity", "Qty"))
        if quantity is None:
            continue

        results.append({
            "symbol": symbol,
            "direction": "short" if quantity < 0 else "long",
            "strike": occ["strike"],
            "expiry": occ["expiry"],
            "underlying": occ["underlying"],
            "quantity": int(quantity),          # signed: negative = short
            "unit_cost": _to_number(cell(row, "Unit Cost", "单位成本", "Avg Cost")),
            "current_price": _to_number(cell(row, "Last Price", "价格", "Price", "Last")),
            "market_value": _to_number(cell(row, "Market Value", "市值", "MktVal")),
            "day_pnl": _to_number(cell(row, "Day Chg $", "$ Day Chg", "当日盈亏")),
            "total_pnl": _to_number(
                cell(row, "$ Gain/Loss", "益损 $", "益损$", "Gain/Loss $", "P&L")),
        })
        seen.add(symbol)

    return results


def stock_snapshot_rows(parsed: list[dict], previous_symbols: Iterable[str]) -> list[dict]:
    """Combine the xlsx snapshot with explicit zero rows for closed positions.

    `positions` is append-only -- no unique constraint on symbol -- and
    `exposure_context.load_portfolio()` reads each symbol's own latest row.
    Appending only what the export still holds is therefore not enough: a name
    sold yesterday keeps its last row as the newest one for that symbol and goes
    on counting toward concentration at its old market value, forever.

    Writing a zero-quantity row instead of deleting history lets
    `compute_exposures()` drop it naturally (it filters zero net value) while
    every historical snapshot stays on disk.
    """
    rows = [dict(row) for row in parsed]
    current = {row["symbol"] for row in rows}
    for symbol in sorted(set(previous_symbols) - current):
        rows.append({
            "symbol": symbol,
            "position_type": "stock",
            "quantity": 0,
            "cost_basis": 0,
            "market_value": 0,
            "unrealized_pnl": 0,
            "unrealized_pnl_pct": 0,
            "description": "已清仓（xlsx 快照中不存在）",
        })
    return rows
